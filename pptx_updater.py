from pptx import Presentation
from pptx.chart.data import CategoryChartData
import io

class PPTXUpdater:
    def __init__(self, template_path):
        """Carrega a apresentação base (Ex: Julho) na memória."""
        self.prs = Presentation(template_path)

    def manter_formatacao_texto(self, shape, novo_texto):
        """Substitui o texto preservando a formatação do shape original."""
        if not shape.has_text_frame: return
        text_frame = shape.text_frame
        
        # Tenta salvar formatação existente
        fonte_nome, fonte_tamanho, fonte_bold, fonte_cor = None, None, None, None
        if text_frame.paragraphs and text_frame.paragraphs[0].runs:
            primeiro_run = text_frame.paragraphs[0].runs[0]
            fonte_nome = primeiro_run.font.name
            fonte_tamanho = primeiro_run.font.size
            fonte_bold = primeiro_run.font.bold
            if hasattr(primeiro_run.font.color, 'rgb') and primeiro_run.font.color.type != 0:
                fonte_cor = primeiro_run.font.color.rgb

        # Limpa e injeta novo texto
        text_frame.clear()
        run = text_frame.paragraphs[0].add_run()
        run.text = str(novo_texto)
        
        # Devolve a formatação
        if fonte_nome: run.font.name = fonte_nome
        if fonte_tamanho: run.font.size = fonte_tamanho
        if fonte_bold is not None: run.font.bold = fonte_bold
        if fonte_cor: run.font.color.rgb = fonte_cor

    def atualizar_texto_por_nome(self, slide_index, shape_name, novo_texto):
        """Busca uma caixa de texto específica pelo nome exato."""
        slide = self.prs.slides[slide_index]
        for shape in slide.shapes:
            if shape.name == shape_name:
                self.manter_formatacao_texto(shape, novo_texto)
                break 

    def obter_arquivo_em_memoria(self):
        """Prepara o arquivo PPTX atualizado para download via Streamlit, sem salvar no servidor."""
        ppt_stream = io.BytesIO()
        self.prs.save(ppt_stream)
        ppt_stream.seek(0)
        return ppt_stream